#!/usr/bin/env python3
"""
TR Change System — Gerador de PPTX para Delivery Follow-up
Lê o JSON exportado pelo template interativo (delivery-follow-up.html)
e gera apresentação executiva no padrão Thomson Reuters.

Uso:
    python generate_pptx.py dados/plataforma-cloud_sprint-14-q2-2026.json
    python generate_pptx.py dados/plataforma-cloud_sprint-14-q2-2026.json saida.pptx

Dependências:
    pip install python-pptx
"""

import json
import sys
from pathlib import Path
from datetime import datetime

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
except ImportError:
    print("Erro: instale python-pptx com 'pip install python-pptx'")
    sys.exit(1)

# ── Paleta Thomson Reuters ──────────────────────────────────────
TR_GREEN_DARK = RGBColor(0x13, 0x3C, 0x2C)
TR_GREEN_MID  = RGBColor(0x1A, 0x5C, 0x3A)
TR_GREEN_LIGHT = RGBColor(0xE8, 0xF0, 0xEB)
WHITE         = RGBColor(0xFF, 0xFF, 0xFF)
GRAY_50       = RGBColor(0xF9, 0xFA, 0xFB)
GRAY_100      = RGBColor(0xF3, 0xF4, 0xF6)
GRAY_200      = RGBColor(0xE5, 0xE7, 0xEB)
GRAY_500      = RGBColor(0x6B, 0x72, 0x80)
GRAY_700      = RGBColor(0x37, 0x41, 0x51)
GRAY_900      = RGBColor(0x11, 0x18, 0x27)
SEM_GREEN     = RGBColor(0x22, 0xC5, 0x5E)
SEM_AMBER     = RGBColor(0xF5, 0x9E, 0x0B)
SEM_RED       = RGBColor(0xEF, 0x44, 0x44)
BLUE          = RGBColor(0x1E, 0x40, 0xAF)

SLIDE_WIDTH  = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

EMOJI_TO_COLOR = {
    "\U0001f7e2": SEM_GREEN,
    "\U0001f7e1": SEM_AMBER,
    "\U0001f534": SEM_RED,
    "\u2b1c":     GRAY_200,
    "On Track":   SEM_GREEN,
    "Atenção":    SEM_AMBER,
    "Bloqueado":  SEM_RED,
}

EMOJI_TO_LABEL = {
    "\U0001f7e2": "On Track",
    "\U0001f7e1": "Atenção",
    "\U0001f534": "Crítico",
    "\u2b1c":     "Futuro",
    "On Track":   "On Track",
    "Atenção":    "Atenção",
    "Bloqueado":  "Bloqueado",
}


# ── Helpers ──────────────────────────────────────────────────────

def cell_val(row, index, default=""):
    """Safely extract value from serialized table cell."""
    if row and index < len(row):
        return row[index].get("value", default)
    return default


def row_is_empty(row):
    """Check if all cell values are empty/zero."""
    return all(
        (cell_val(row, i) or "").strip() in ("", "0")
        for i in range(len(row))
    )


def filter_rows(rows):
    """Remove entirely empty rows."""
    if not rows:
        return []
    return [r for r in rows if not row_is_empty(r)]


def status_color(raw):
    """Map emoji or text status to RGBColor."""
    return EMOJI_TO_COLOR.get(raw, GRAY_500)


def status_label(raw):
    """Map emoji to readable label."""
    return EMOJI_TO_LABEL.get(raw, raw)


def format_date(val):
    """Try to format ISO date to dd/mm/yyyy."""
    if not val:
        return ""
    try:
        return datetime.fromisoformat(val).strftime("%d/%m/%Y")
    except (ValueError, TypeError):
        return val


# ── Slide building blocks ────────────────────────────────────────

def add_header(slide, label, title, subtitle=""):
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), SLIDE_WIDTH, Inches(1.15))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TR_GREEN_DARK
    bar.line.fill.background()

    tf = bar.text_frame
    tf.word_wrap = True
    tf.margin_top = Pt(14)
    tf.margin_left = Pt(40)

    p0 = tf.paragraphs[0]
    p0.text = label.upper()
    p0.font.size = Pt(8)
    p0.font.color.rgb = WHITE
    p0.font.bold = True
    p0.space_after = Pt(2)

    p1 = tf.add_paragraph()
    p1.text = title
    p1.font.size = Pt(20)
    p1.font.color.rgb = WHITE
    p1.font.bold = True

    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(10)
        p2.font.color.rgb = WHITE
        p2.space_before = Pt(2)


def add_footer(slide, extra=""):
    tb = slide.shapes.add_textbox(Inches(0), Inches(7.05), SLIDE_WIDTH, Inches(0.35))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    txt = "Thomson Reuters · Delivery Follow-up · Confidencial"
    if extra:
        txt += f"   |   {extra}"
    p.text = txt
    p.font.size = Pt(7)
    p.font.color.rgb = GRAY_500
    p.alignment = PP_ALIGN.CENTER


def add_section_title(slide, left, top, text, color=TR_GREEN_DARK):
    tb = slide.shapes.add_textbox(left, top, Inches(11), Inches(0.35))
    p = tb.text_frame.paragraphs[0]
    p.text = text.upper()
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = color
    p.space_after = Pt(4)


def add_kpi_box(slide, left, top, width, value, label, color=GRAY_900):
    box = slide.shapes.add_shape(1, left, top, width, Inches(1.05))
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = GRAY_200
    box.line.width = Pt(1)

    tf = box.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    pv = tf.paragraphs[0]
    pv.text = str(value)
    pv.font.size = Pt(32)
    pv.font.bold = True
    pv.font.color.rgb = color

    pl = tf.add_paragraph()
    pl.text = label.upper()
    pl.font.size = Pt(7)
    pl.font.color.rgb = GRAY_500
    pl.alignment = PP_ALIGN.CENTER


def build_table(slide, left, top, width, col_widths, headers, rows_data):
    """Build a styled pptx table. Returns the table shape."""
    n_rows = len(rows_data) + 1
    n_cols = len(headers)
    tbl_shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, Inches(0.35 * n_rows))
    tbl = tbl_shape.table

    for i, w in enumerate(col_widths):
        tbl.columns[i].width = w

    for i, h in enumerate(headers):
        cell = tbl.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = TR_GREEN_DARK
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(8)
        p.font.bold = True
        p.font.color.rgb = WHITE
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    for r, row in enumerate(rows_data):
        for c in range(n_cols):
            cell = tbl.cell(r + 1, c)
            val = row[c] if c < len(row) else ""
            cell.text = str(val)
            bg = WHITE if r % 2 == 0 else GRAY_50
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(9)
            p.font.color.rgb = GRAY_900
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    return tbl_shape


def style_status_cell(tbl, row_idx, col_idx, raw_status):
    """Apply colored fill to a status cell."""
    cell = tbl.table.cell(row_idx, col_idx)
    color = status_color(raw_status)
    cell.fill.solid()
    cell.fill.fore_color.rgb = color
    p = cell.text_frame.paragraphs[0]
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.font.size = Pt(8)
    cell.text = status_label(raw_status)


# ── Slide builders ───────────────────────────────────────────────

def slide_dashboard(prs, data):
    """Slide 1: Dashboard Consolidado — KPIs + Progresso."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    meta = data.get("meta", {})
    produto = meta.get("produto", "[Programa / Produto]")
    periodo = meta.get("periodo", "[Período]")
    dt = format_date(meta.get("data", ""))
    resp = meta.get("responsavel", "")

    add_header(
        slide,
        "TR Change System · Delivery Follow-up",
        f"{produto} · {periodo}",
        f"{dt}   ·   {resp}   ·   Confidencial" if dt else ""
    )

    andamento_count = len(filter_rows(data.get("tbl-andamento", [])))
    PURPLE = RGBColor(0x6D, 0x28, 0xD9)
    kpi_defs = [
        ("kpi-assumidas",  "Entregas Assumidas",       BLUE),
        ("kpi-realizadas", "Entregas Realizadas",       SEM_GREEN),
        (None,             "Iniciativas em Andamento",  PURPLE, str(andamento_count)),
        ("kpi-bloqueadas", "Bloqueadas",                SEM_RED),
        ("kpi-atrasadas",  "Atrasadas",                 SEM_AMBER),
    ]
    kpi_w = Inches(2.3)
    gap = Inches(0.2)
    x0 = Inches(0.5)
    y_kpi = Inches(1.55)

    for i, item in enumerate(kpi_defs):
        key, label, color = item[0], item[1], item[2]
        val = item[3] if len(item) > 3 else data.get(key, "0")
        add_kpi_box(slide, x0 + i * (kpi_w + gap), y_kpi, kpi_w, val, label, color)

    y_prog = Inches(2.95)
    add_section_title(slide, x0, y_prog, "Progresso")

    prog_items = [
        ("prog-conclusao", "% Conclusão"),
        ("prog-velocity",  "Velocity"),
        ("prog-backlog",   "Features no Backlog"),
    ]
    pw = Inches(3.6)
    for i, (key, label) in enumerate(prog_items):
        val = data.get(key, "—")
        bx = slide.shapes.add_shape(1, x0 + i * (pw + Inches(0.2)), y_prog + Inches(0.4), pw, Inches(0.65))
        bx.fill.solid()
        bx.fill.fore_color.rgb = TR_GREEN_LIGHT
        bx.line.fill.background()
        tf = bx.text_frame
        tf.word_wrap = True
        tf.margin_left = Pt(12)
        tf.paragraphs[0].text = str(val) if val else "—"
        tf.paragraphs[0].font.size = Pt(18)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = TR_GREEN_DARK
        p2 = tf.add_paragraph()
        p2.text = label
        p2.font.size = Pt(8)
        p2.font.color.rgb = GRAY_500

    add_footer(slide)


def slide_entregas(prs, data):
    """Slide 2: Principais Entregas do Período."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Delivery Follow-up", "2. Principais Entregas do Período")

    x0 = Inches(0.5)
    y = Inches(1.45)

    concluidas = filter_rows(data.get("tbl-concluidas", []))
    if concluidas:
        add_section_title(slide, x0, y, "Entregas Concluídas", SEM_GREEN)
        y += Inches(0.35)
        rows = [[cell_val(r, 0), cell_val(r, 1), format_date(cell_val(r, 2))] for r in concluidas]
        t = build_table(slide, x0, y, Inches(12.3),
                        [Inches(2.5), Inches(7.8), Inches(2.0)],
                        ["Frente", "Entrega", "Data"], rows)
        y += Inches(0.35 * (len(rows) + 1) + 0.2)

    andamento = filter_rows(data.get("tbl-andamento", []))
    if andamento:
        add_section_title(slide, x0, y, "Em Andamento", SEM_AMBER)
        y += Inches(0.35)
        rows = []
        for r in andamento:
            rows.append([
                cell_val(r, 0), cell_val(r, 1),
                cell_val(r, 2), cell_val(r, 3),
                format_date(cell_val(r, 4)),
                status_label(cell_val(r, 5))
            ])
        t = build_table(slide, x0, y, Inches(12.3),
                        [Inches(2.0), Inches(4.0), Inches(1.2), Inches(1.2), Inches(2.0), Inches(1.9)],
                        ["Frente", "Entrega", "Previsto", "Realizado", "Target", "Status"], rows)
        for ri, r in enumerate(andamento):
            style_status_cell(t, ri + 1, 5, cell_val(r, 5))
        y += Inches(0.35 * (len(rows) + 1) + 0.2)

    atencao = filter_rows(data.get("tbl-atencao", []))
    if atencao:
        add_section_title(slide, x0, y, "Itens em Atenção", SEM_RED)
        y += Inches(0.35)
        rows = [[cell_val(r, 0), cell_val(r, 1), cell_val(r, 2), cell_val(r, 3)] for r in atencao]
        build_table(slide, x0, y, Inches(12.3),
                    [Inches(2.5), Inches(4.3), Inches(3.5), Inches(2.0)],
                    ["Frente", "Item", "Motivo", "Dias Atraso"], rows)

    add_footer(slide)


def slide_riscos(prs, data):
    """Slide 3: Riscos e Plano de Ação."""
    riscos = filter_rows(data.get("tbl-riscos", []))
    if not riscos:
        return

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Delivery Follow-up", "3. Riscos e Plano de Ação")

    rows = []
    for r in riscos:
        rows.append([
            status_label(cell_val(r, 0)),
            cell_val(r, 1), cell_val(r, 2),
            cell_val(r, 3), cell_val(r, 4),
            format_date(cell_val(r, 5))
        ])

    t = build_table(
        slide, Inches(0.5), Inches(1.55), Inches(12.3),
        [Inches(1.3), Inches(3.0), Inches(2.0), Inches(3.0), Inches(1.5), Inches(1.5)],
        ["Sev.", "Risco", "Área Afetada", "Plano de Ação", "Owner", "Prazo"],
        rows
    )
    for ri, r in enumerate(riscos):
        style_status_cell(t, ri + 1, 0, cell_val(r, 0))

    add_footer(slide)


def slide_roadmap(prs, data):
    """Slide 4: Roadmap — split into Outcomes and Projeto tables."""
    milestones = filter_rows(data.get("tbl-milestones", []))
    if not milestones:
        return

    outcomes = [r for r in milestones if cell_val(r, 0) != "Projeto"]
    projetos = [r for r in milestones if cell_val(r, 0) == "Projeto"]

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Delivery Follow-up", "4. Roadmap")

    y = Inches(1.45)
    col_w = [Inches(2.5), Inches(3.0), Inches(1.5), Inches(1.5), Inches(1.8), Inches(2.0)]

    if outcomes:
        add_section_title(slide, Inches(0.5), y, "Outcomes", TR_GREEN_MID)
        y += Inches(0.35)
        rows = []
        for r in outcomes:
            rows.append([
                cell_val(r, 1), cell_val(r, 2), cell_val(r, 3),
                cell_val(r, 4), format_date(cell_val(r, 5)), status_label(cell_val(r, 6))
            ])
        t = build_table(slide, Inches(0.5), y, Inches(12.3), col_w,
                        ["Iniciativa / Frente", "Resultado esperado", "Baseline", "Valor Atual", "Previsão", "Status"], rows)
        for ri, r in enumerate(outcomes):
            style_status_cell(t, ri + 1, 5, cell_val(r, 6))
        y += Inches(0.35 * (len(rows) + 1) + 0.4)

    if projetos:
        add_section_title(slide, Inches(0.5), y, "Projeto", BLUE)
        y += Inches(0.35)
        rows = []
        for r in projetos:
            rows.append([
                cell_val(r, 1), cell_val(r, 2), cell_val(r, 3),
                cell_val(r, 4), format_date(cell_val(r, 5)), status_label(cell_val(r, 6))
            ])
        t = build_table(slide, Inches(0.5), y, Inches(12.3), col_w,
                        ["Iniciativa / Frente", "Resultado esperado", "Previsto", "Realizado", "Previsão", "Status"], rows)
        for ri, r in enumerate(projetos):
            style_status_cell(t, ri + 1, 5, cell_val(r, 6))

    add_footer(slide)


def slide_enablers(prs, data):
    """Slide 5: Enablers."""
    enablers = filter_rows(data.get("tbl-enablers", []))
    if not enablers:
        return

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Delivery Follow-up", "5. Enablers")

    rows = []
    for r in enablers:
        prog = cell_val(r, 2)
        prog_str = f"{prog}%" if prog and prog != "0" else "—"
        rows.append([
            cell_val(r, 0), cell_val(r, 1),
            prog_str, format_date(cell_val(r, 3)),
            cell_val(r, 4), cell_val(r, 5)
        ])

    t = build_table(
        slide, Inches(0.5), Inches(1.55), Inches(12.3),
        [Inches(3.0), Inches(1.5), Inches(1.5), Inches(1.8), Inches(2.0), Inches(2.5)],
        ["Iniciativa", "Status", "Progresso", "Previsão", "Owner", "Observações"],
        rows
    )
    for ri, r in enumerate(enablers):
        raw = cell_val(r, 1)
        if raw in EMOJI_TO_COLOR or raw in ("On Track", "Atenção", "Bloqueado"):
            style_status_cell(t, ri + 1, 1, raw)

    add_footer(slide)


def slide_decisoes(prs, data):
    """Slide 6: Decisões Necessárias."""
    decisoes = filter_rows(data.get("tbl-decisoes", []))
    if not decisoes:
        return

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "Delivery Follow-up", "6. Decisões Necessárias")

    tb = slide.shapes.add_textbox(Inches(0.5), Inches(1.35), Inches(10), Inches(0.3))
    p = tb.text_frame.paragraphs[0]
    p.text = "Itens que requerem decisão da liderança neste follow-up."
    p.font.size = Pt(10)
    p.font.color.rgb = GRAY_500

    rows = []
    for i, r in enumerate(decisoes):
        rows.append([
            str(i + 1),
            cell_val(r, 0), cell_val(r, 1),
            cell_val(r, 2), cell_val(r, 3),
            format_date(cell_val(r, 4))
        ])

    t = build_table(
        slide, Inches(0.5), Inches(1.75), Inches(12.3),
        [Inches(0.5), Inches(2.8), Inches(2.8), Inches(2.2), Inches(2.5), Inches(1.5)],
        ["#", "Decisão", "Contexto", "Opções", "Recomendação", "Prazo"],
        rows
    )

    for ri in range(len(rows)):
        cell = t.table.cell(ri + 1, 4)
        cell.text_frame.paragraphs[0].font.bold = True

    # amber left-border effect: color the # column cells
    for ri in range(len(rows)):
        cell = t.table.cell(ri + 1, 0)
        cell.fill.solid()
        cell.fill.fore_color.rgb = SEM_AMBER
        cell.text_frame.paragraphs[0].font.color.rgb = WHITE
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    add_footer(slide)


# ── Main ─────────────────────────────────────────────────────────

def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(0)

    json_path = Path(sys.argv[1])
    if not json_path.exists():
        print(f"Erro: arquivo '{json_path}' não encontrado.")
        sys.exit(1)

    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    output = sys.argv[2] if len(sys.argv) > 2 else json_path.with_suffix(".pptx")

    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    slide_dashboard(prs, data)
    slide_entregas(prs, data)
    slide_riscos(prs, data)
    slide_roadmap(prs, data)
    slide_enablers(prs, data)
    slide_decisoes(prs, data)

    prs.save(str(output))
    meta = data.get("meta", {})
    print(f"PPTX gerado: {output}")
    print(f"  {meta.get('produto', '')} · {meta.get('periodo', '')} · {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
